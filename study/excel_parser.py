from __future__ import annotations
import os
import re
import html
import logging
from datetime import datetime, date, time as dt_time
from typing import List, Dict, Any, Tuple, Optional, Set

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell

from study.pdf_parser import detect_language

logger = logging.getLogger("study-excel-parser")

def inspect_excel(file_path_or_bytes: str | bytes) -> Dict[str, Any]:
    """
    Quickly inspects an Excel workbook to extract metadata (sheet names, total sheets).
    """
    import io
    if isinstance(file_path_or_bytes, bytes):
        wb = openpyxl.load_workbook(io.BytesIO(file_path_or_bytes), read_only=True, data_only=True)
    else:
        wb = openpyxl.load_workbook(file_path_or_bytes, read_only=True, data_only=True)
    
    sheet_names = wb.sheetnames
    wb.close()
    return {
        "sheets_count": len(sheet_names),
        "sheets": sheet_names
    }

def _format_cell_value(cell: Cell | Any) -> str:
    """
    Formats an openpyxl cell value into a clean, safe HTML string.
    Preserves newlines as <br/>, handles dates, floats, and escapes special HTML entities.
    """
    val = getattr(cell, "value", None) if hasattr(cell, "value") else cell
    if val is None:
        return ""
    
    if isinstance(val, (datetime, date)):
        if isinstance(val, datetime):
            formatted = val.strftime("%Y-%m-%d %H:%M:%S" if (val.hour or val.minute or val.second) else "%Y-%m-%d")
        else:
            formatted = val.strftime("%Y-%m-%d")
        return html.escape(formatted)
    
    if isinstance(val, dt_time):
        return html.escape(val.strftime("%H:%M:%S"))

    if isinstance(val, float):
        # Format float cleanly, avoid long trailing precision like 12.0000000000001
        formatted = f"{val:g}"
        return html.escape(formatted)
    
    if isinstance(val, bool):
        return "TRUE" if val else "FALSE"

    raw_text = str(val).strip()
    if not raw_text:
        return ""

    # Escape HTML special chars and preserve cell linebreaks
    escaped = html.escape(raw_text)
    return escaped.replace("\r\n", "<br/>").replace("\r", "<br/>").replace("\n", "<br/>")

def _get_sheet_bounding_box(ws: Worksheet) -> Tuple[int, int, int, int]:
    """
    Finds the actual populated bounding box (min_row, max_row, min_col, max_col)
    to avoid parsing thousands of ghost empty rows or columns.
    """
    min_r = ws.max_row or 1
    max_r = 1
    min_c = ws.max_column or 1
    max_c = 1
    has_data = False

    # Also consider merged cell ranges for bounds
    for rng in ws.merged_cells.ranges:
        has_data = True
        min_r = min(min_r, rng.min_row)
        max_r = max(max_r, rng.max_row)
        min_c = min(min_c, rng.min_col)
        max_c = max(max_c, rng.max_col)

    # Scan rows
    raw_max_r = min(ws.max_row or 1, 2000)  # Safety limit
    raw_max_c = min(ws.max_column or 1, 100)

    for r in range(1, raw_max_r + 1):
        for c in range(1, raw_max_c + 1):
            val = ws.cell(r, c).value
            if val is not None and str(val).strip():
                has_data = True
                min_r = min(min_r, r)
                max_r = max(max_r, r)
                min_c = min(min_c, c)
                max_c = max(max_c, c)

    if not has_data:
        return 1, 1, 1, 1

    return min_r, max_r, min_c, max_c

def sheet_to_html_table(ws: Worksheet) -> Tuple[str, str]:
    """
    Converts a worksheet to a well-structured, responsive HTML table.
    Properly handles cell merging (rowspan, colspan) and cell alignments.
    Returns: (html_table_string, extracted_plain_text)
    """
    # 1. Resolve merged cell spans and skip sets
    spans: Dict[Tuple[int, int], Tuple[int, int]] = {}
    skip_cells: Set[Tuple[int, int]] = set()

    for rng in ws.merged_cells.ranges:
        r_span = rng.max_row - rng.min_row + 1
        c_span = rng.max_col - rng.min_col + 1
        spans[(rng.min_row, rng.min_col)] = (r_span, c_span)

        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    skip_cells.add((r, c))

    # 2. Determine non-empty bounding box
    min_r, max_r, min_c, max_c = _get_sheet_bounding_box(ws)

    # Check if empty sheet
    if min_r == max_r and min_c == max_c:
        first_val = ws.cell(min_r, min_c).value
        if first_val is None or not str(first_val).strip():
            return '<div class="excel-empty-sheet text-muted-foreground italic py-4">（此工作表无内容）</div>', ""

    html_lines = [
        '<div class="excel-table-container overflow-x-auto my-3 border border-border/80 rounded-lg shadow-sm">',
        '  <table class="excel-table min-w-full border-collapse text-sm">'
    ]
    
    extracted_text_pieces: List[str] = []

    for r in range(min_r, max_r + 1):
        # Check if entire row is empty and not merged
        has_row_content = False
        for c in range(min_c, max_c + 1):
            coord = (r, c)
            if coord in spans or coord in skip_cells:
                has_row_content = True
                break
            val = ws.cell(r, c).value
            if val is not None and str(val).strip():
                has_row_content = True
                break

        if not has_row_content:
            continue

        html_lines.append("    <tr class=\"hover:bg-muted/30 transition-colors\">")
        
        is_first_row = (r == min_r)

        for c in range(min_c, max_c + 1):
            coord = (r, c)
            if coord in skip_cells:
                # Part of a merged span, skipped in HTML
                continue

            cell = ws.cell(r, c)
            cell_html = _format_cell_value(cell)
            if cell.value is not None and str(cell.value).strip():
                extracted_text_pieces.append(str(cell.value).strip())

            attrs: List[str] = []
            
            # Merged attributes
            if coord in spans:
                rowspan, colspan = spans[coord]
                if rowspan > 1:
                    attrs.append(f'rowspan="{rowspan}"')
                if colspan > 1:
                    attrs.append(f'colspan="{colspan}"')

            # Formatting & alignment styles
            styles: List[str] = []
            if cell.alignment:
                if cell.alignment.horizontal in ("left", "center", "right"):
                    styles.append(f"text-align: {cell.alignment.horizontal}")
                if cell.alignment.vertical in ("top", "center", "bottom"):
                    styles.append(f"vertical-align: {cell.alignment.vertical}")

            is_bold = bool(cell.font and cell.font.bold)
            tag = "th" if is_first_row else "td"
            
            # Base Tailwind classes
            if tag == "th":
                cls = "border border-border/70 bg-muted/60 px-3 py-2 text-left font-semibold text-foreground text-xs select-text"
            else:
                cls = "border border-border/60 px-3 py-2 text-foreground/90 text-xs align-top select-text"

            if is_bold and tag != "th":
                cls += " font-semibold"

            attrs.append(f'class="{cls}"')
            if styles:
                attrs.append(f'style="{"; ".join(styles)}"')

            content = cell_html if cell_html else "&nbsp;"
            attr_str = " " + " ".join(attrs) if attrs else ""
            html_lines.append(f"      <{tag}{attr_str}>{content}</{tag}>")

        html_lines.append("    </tr>")

    html_lines.append("  </table>")
    html_lines.append("</div>")

    return "\n".join(html_lines), " ".join(extracted_text_pieces)


def extract_excel_content(
    doc_id: str,
    doc_dir: str,
    filename: str = "",
    progress_callback: Optional[Any] = None
) -> Tuple[str, List[Dict[str, Any]]]:
    """
    Main extraction entrypoint for Excel (.xlsx, .xlsm, .xltx, etc.):
    - Each Worksheet corresponds to one Chapter.
    - Sheet contents are converted to styled, merged-cell supporting HTML tables.
    - Detects primary document language across all sheets.
    """
    # Locate original excel file
    file_path = None
    for candidate in os.listdir(doc_dir):
        if candidate.startswith("original."):
            ext = os.path.splitext(candidate)[1].lower()
            if ext in (".xlsx", ".xlsm", ".xltx", ".xltm", ".xls"):
                file_path = os.path.join(doc_dir, candidate)
                break

    if not file_path or not os.path.exists(file_path):
        raise FileNotFoundError(f"Original Excel file for document {doc_id} not found in {doc_dir}")

    logger.info(f"Loading Excel workbook: {file_path}")
    
    # Load workbook with data_only=True so formulas return calculated values
    wb = openpyxl.load_workbook(file_path, data_only=True)
    all_sheetnames = wb.sheetnames
    total_sheets = len(all_sheetnames)

    chapters: List[Dict[str, Any]] = []
    all_extracted_text: List[str] = []

    for idx, sheet_name in enumerate(all_sheetnames, start=1):
        if progress_callback:
            progress_callback(idx - 1, total_sheets, f"正在解析工作表: {sheet_name}")

        ws = wb[sheet_name]
        
        # Check sheet visibility
        state = getattr(ws, "sheet_state", "visible")
        sheet_display_title = sheet_name
        if state == "hidden":
            sheet_display_title += " (已隐藏)"

        table_html, sheet_text = sheet_to_html_table(ws)
        if sheet_text:
            all_extracted_text.append(sheet_text)

        chapter_id = f"ch_{idx:03d}"
        para_id = f"p_{idx:03d}_0001"

        paragraph = {
            "id": para_id,
            "chapter_id": chapter_id,
            "order": 1,
            "type": "table",
            "block_type": "table",
            "source_text": table_html,
            "english": table_html,
            "chinese": "",
            "translated_text": "",
            "status": "pending",
            "is_heading": False,
            "page_num": idx,
            "notes": []
        }

        chapters.append({
            "chapter_id": chapter_id,
            "title": sheet_display_title,
            "paragraphs": [paragraph]
        })

    wb.close()

    if progress_callback:
        progress_callback(total_sheets, total_sheets, "工作表解析完成")

    # Detect language across entire workbook content
    combined_sample = " ".join(all_extracted_text)
    if combined_sample.strip():
        lang = detect_language(combined_sample[:3000])
    else:
        lang = "zh"

    # If Chinese, mark all paragraphs as completed without translation needed
    if lang == "zh":
        for ch in chapters:
            for p in ch["paragraphs"]:
                p["chinese"] = p["source_text"]
                p["translated_text"] = p["source_text"]
                p["status"] = "completed"

    if not chapters:
        chapters.append({
            "chapter_id": "ch_001",
            "title": "Sheet1",
            "paragraphs": []
        })

    logger.info(f"Excel extraction completed for doc {doc_id}: {len(chapters)} chapters, detected language: {lang}")
    return lang, chapters
