import io
import os
import re
import logging
from typing import List, Dict, Any, Tuple, Optional
try:
    from PIL import Image, ImageDraw, ImageFont
    HAS_PIL = True
except ImportError:
    HAS_PIL = False
    Image = None
    ImageDraw = None
    ImageFont = None

logger = logging.getLogger("doc-converter")


# Try to import PyMuPDF (fitz)
try:
    import fitz  # PyMuPDF
    HAS_PYMUPDF = True
except ImportError:
    HAS_PYMUPDF = False
    logger.warning("PyMuPDF (fitz) is not installed. PDF conversion will be limited.")

# Try to import python-pptx
try:
    from pptx import Presentation
    HAS_PPTX = True
except ImportError:
    HAS_PPTX = False

# Try to import python-docx
try:
    from docx import Document as DocxDocument
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

# Try to import win32com on Windows if available
try:
    import win32com.client
    import pythoncom
    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

def parse_page_range(range_str: str, max_pages: int) -> List[int]:
    """
    Parses a page range string like '1-3, 5, 7-9' or 'all' into a list of 1-based page numbers.
    """
    if not range_str or range_str.strip().lower() in ("all", "全部", "*"):
        return list(range(1, max_pages + 1))
    
    selected_pages = set()
    parts = [p.strip() for p in range_str.split(",") if p.strip()]
    for part in parts:
        if "-" in part:
            sub = part.split("-")
            if len(sub) == 2:
                try:
                    start = int(sub[0].strip())
                    end = int(sub[1].strip())
                    for p in range(min(start, end), max(start, end) + 1):
                        if 1 <= p <= max_pages:
                            selected_pages.add(p)
                except ValueError:
                    continue
        else:
            try:
                p = int(part)
                if 1 <= p <= max_pages:
                    selected_pages.add(p)
            except ValueError:
                continue
                
    result = sorted(list(selected_pages))
    if not result:
        return list(range(1, min(max_pages, 5) + 1))
    return result

def convert_pdf_to_images(file_bytes: bytes, pages: Optional[List[int]] = None, dpi: int = 150) -> List[Tuple[int, bytes]]:
    """
    Converts PDF pages to PNG byte images using PyMuPDF.
    Returns list of (page_num, png_bytes).
    """
    if not HAS_PYMUPDF:
        raise RuntimeError("PyMuPDF (fitz) is required for PDF conversion. Please install with: pip install pymupdf")
    
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    total_pages = len(doc)
    if pages is None:
        target_pages = list(range(1, total_pages + 1))
    else:
        target_pages = [p for p in pages if 1 <= p <= total_pages]
        
    result = []
    zoom = dpi / 72.0
    matrix = fitz.Matrix(zoom, zoom)
    
    for page_num in target_pages:
        page = doc.load_page(page_num - 1)
        pix = page.get_pixmap(matrix=matrix, alpha=False)
        png_bytes = pix.tobytes("png")
        result.append((page_num, png_bytes))
        
    doc.close()
    return result

def get_pdf_page_count(file_bytes: bytes) -> int:
    if not HAS_PYMUPDF:
        return 1
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    count = len(doc)
    doc.close()
    return count

def render_pptx_slides_fallback(file_bytes: bytes, pages: Optional[List[int]] = None) -> List[Tuple[int, bytes]]:
    """
    Renders PPTX slides to images using python-pptx and Pillow.
    Extracts text and embedded images to create visual slide approximations.
    """
    if not HAS_PPTX:
        raise RuntimeError("python-pptx is required for PPTX conversion. Please install with: pip install python-pptx")
    
    prs = Presentation(io.BytesIO(file_bytes))
    total_slides = len(prs.slides)
    if pages is None:
        target_pages = list(range(1, total_slides + 1))
    else:
        target_pages = [p for p in pages if 1 <= p <= total_slides]
        
    slide_width = prs.slide_width
    slide_height = prs.slide_height
    # Aspect ratio mapping to approx 1280x720
    img_w = 1280
    img_h = int(img_w * (slide_height / slide_width)) if slide_width else 720
    
    result = []
    for page_num in target_pages:
        slide = prs.slides[page_num - 1]
        img = Image.new("RGB", (img_w, img_h), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        
        # Draw header bar
        draw.rectangle([0, 0, img_w, 40], fill=(240, 243, 246))
        draw.text((20, 10), f"Slide {page_num} / {total_slides}", fill=(100, 116, 139))
        
        y_offset = 60
        for shape in slide.shapes:
            # Handle images inside slide
            if shape.shape_type == 13:  # Picture
                try:
                    img_blob = shape.image.blob
                    pil_shape_img = Image.open(io.BytesIO(img_blob))
                    # Scale down if needed
                    pil_shape_img.thumbnail((400, 300))
                    img.paste(pil_shape_img, (img_w - 420, y_offset))
                except Exception:
                    pass
                    
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        # Draw text lines
                        draw.text((40, y_offset), text[:120], fill=(30, 41, 59))
                        y_offset += 32
                        if y_offset > img_h - 40:
                            break
                            
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        result.append((page_num, buf.getvalue()))
        
    return result

def get_pptx_slide_count(file_bytes: bytes) -> int:
    if not HAS_PPTX:
        return 1
    try:
        prs = Presentation(io.BytesIO(file_bytes))
        return len(prs.slides)
    except Exception:
        return 1

def render_docx_pages_fallback(file_bytes: bytes, pages: Optional[List[int]] = None) -> List[Tuple[int, bytes]]:
    """
    Renders DOCX document pages to images using python-docx and Pillow.
    """
    if not HAS_DOCX:
        raise RuntimeError("python-docx is required for DOCX conversion. Please install with: pip install python-docx")
        
    doc = DocxDocument(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    
    # Estimate pages (around 20-25 paragraphs per page)
    lines_per_page = 22
    chunks = []
    for i in range(0, max(len(paragraphs), 1), lines_per_page):
        chunks.append(paragraphs[i:i+lines_per_page])
        
    total_pages = len(chunks) if chunks else 1
    if pages is None:
        target_pages = list(range(1, total_pages + 1))
    else:
        target_pages = [p for p in pages if 1 <= p <= total_pages]
        
    img_w, img_h = 1000, 1400  # Standard A4-ish ratio
    result = []
    
    for page_num in target_pages:
        page_paras = chunks[page_num - 1] if page_num <= len(chunks) else []
        img = Image.new("RGB", (img_w, img_h), color=(255, 255, 255))
        draw = ImageDraw.Draw(img)
        
        # Border and header
        draw.rectangle([0, 0, img_w, 50], fill=(248, 250, 252))
        draw.text((40, 15), f"Page {page_num} of {total_pages}", fill=(100, 116, 139))
        draw.line([0, 50, img_w, 50], fill=(226, 232, 240), width=2)
        
        y = 80
        for para in page_paras:
            # Wrap long text into multiple lines
            words = para
            while len(words) > 50:
                draw.text((60, y), words[:50], fill=(15, 23, 42))
                y += 28
                words = words[50:]
                if y > img_h - 60:
                    break
            if y <= img_h - 60 and words:
                draw.text((60, y), words, fill=(15, 23, 42))
                y += 38
            if y > img_h - 60:
                break
                
        buf = io.BytesIO()
        img.save(buf, format="PNG")
        result.append((page_num, buf.getvalue()))
        
    return result

def get_docx_page_count(file_bytes: bytes) -> int:
    if not HAS_DOCX:
        return 1
    try:
        doc = DocxDocument(io.BytesIO(file_bytes))
        paras = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        return max(1, (len(paras) + 21) // 22)
    except Exception:
        return 1

def inspect_document(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """
    Analyzes uploaded document and returns format, total page count, and metadata.
    """
    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".pdf":
        count = get_pdf_page_count(file_bytes)
        return {"file_type": "pdf", "total_pages": count, "filename": filename}
    elif ext in (".pptx", ".ppt"):
        count = get_pptx_slide_count(file_bytes)
        return {"file_type": "pptx", "total_pages": count, "filename": filename}
    elif ext in (".docx", ".doc"):
        count = get_docx_page_count(file_bytes)
        return {"file_type": "docx", "total_pages": count, "filename": filename}
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"):
        return {"file_type": "image", "total_pages": 1, "filename": filename}
    elif ext == ".txt":
        try:
            line_count = len(file_bytes.splitlines())
            pages = max(1, (line_count + 49) // 50)
        except Exception:
            pages = 1
        return {"file_type": "txt", "total_pages": pages, "filename": filename}
    elif ext in (".md", ".markdown"):
        try:
            line_count = len(file_bytes.splitlines())
            pages = max(1, (line_count + 49) // 50)
        except Exception:
            pages = 1
        return {"file_type": "md", "total_pages": pages, "filename": filename}
    else:
        # Fallback text/binary file
        return {"file_type": "unknown", "total_pages": 1, "filename": filename}

def convert_document_to_page_images(file_bytes: bytes, filename: str, pages_str: Optional[str] = None, dpi: int = 150) -> List[Tuple[int, bytes]]:
    """
    Main dispatch function to convert any supported document into a list of (page_num, png_bytes).
    """
    ext = os.path.splitext(filename)[1].lower()
    info = inspect_document(file_bytes, filename)
    total_pages = info["total_pages"]
    
    selected_pages = parse_page_range(pages_str or "all", total_pages)
    
    if ext == ".pdf":
        return convert_pdf_to_images(file_bytes, selected_pages, dpi=dpi)
    elif ext in (".pptx", ".ppt"):
        return render_pptx_slides_fallback(file_bytes, selected_pages)
    elif ext in (".docx", ".doc"):
        return render_docx_pages_fallback(file_bytes, selected_pages)
    elif ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".gif"):
        return [(1, file_bytes)]
    else:
        raise ValueError(f"Unsupported document format: {ext}")

def generate_document_thumbnails(file_bytes: bytes, filename: str, start_page: int = 1, max_pages: int = 24) -> Tuple[int, List[Dict[str, Any]]]:
    """
    Generates lightweight base64 thumbnail previews for visual page selection.
    Returns (total_pages, thumbnails_list).
    """
    import base64
    info = inspect_document(file_bytes, filename)
    total_pages = info["total_pages"]
    
    if total_pages <= 0 or start_page > total_pages or start_page < 1:
        return total_pages, []
        
    end_page = min(total_pages, start_page + max_pages - 1)
    target_range_str = f"{start_page}-{end_page}"
    
    # Render with lower DPI for speed and bandwidth efficiency
    images = convert_document_to_page_images(file_bytes, filename, pages_str=target_range_str, dpi=60)
    
    thumbnails = []
    for page_num, img_bytes in images:
        try:
            pil_img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
            pil_img.thumbnail((260, 260))
            thumb_io = io.BytesIO()
            pil_img.save(thumb_io, format="JPEG", quality=70)
            b64 = base64.b64encode(thumb_io.getvalue()).decode("utf-8")
            thumbnails.append({
                "page": page_num,
                "thumbnail": f"data:image/jpeg;base64,{b64}"
            })
        except Exception as e:
            logger.error(f"Thumbnail generation error for page {page_num}: {e}")
            
    return total_pages, thumbnails


