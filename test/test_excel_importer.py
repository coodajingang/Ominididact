import os
import io
import shutil
import tempfile
import unittest
import openpyxl
from datetime import datetime

from study.excel_parser import inspect_excel, sheet_to_html_table, extract_excel_content
import study.service as study_service

class TestExcelImporter(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp()
        self.excel_path = os.path.join(self.test_dir, "test_sample.xlsx")
        
        # Create a test Excel workbook
        wb = openpyxl.Workbook()
        
        # Sheet 1: Chinese with merged cells
        ws1 = wb.active
        ws1.title = "财务统计表"
        ws1["A1"] = "部门汇总"
        ws1.merge_cells("A1:B1")  # Colspan 2
        ws1["C1"] = "日期"
        ws1["D1"] = datetime(2026, 10, 21)
        
        ws1["A2"] = "技术部"
        ws1.merge_cells("A2:A3")  # Rowspan 2
        ws1["B2"] = "研发支出"
        ws1["C2"] = 12000.5
        
        ws1["B3"] = "运维支出"
        ws1["C3"] = 4500.0
        
        # Sheet 2: English sheet
        ws2 = wb.create_sheet(title="Project Milestones")
        ws2["A1"] = "Phase"
        ws2["B1"] = "Description"
        ws2["C1"] = "Status"
        ws2["A2"] = "Phase 1"
        ws2["B2"] = "Requirements and architecture design"
        ws2["C2"] = "Done"
        ws2["A3"] = "Phase 2"
        ws2["B3"] = "Implementation and unit testing"
        ws2["C3"] = "In Progress"

        # Sheet 3: Empty sheet
        ws3 = wb.create_sheet(title="Empty Notes")

        wb.save(self.excel_path)
        wb.close()

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_inspect_excel(self):
        info = inspect_excel(self.excel_path)
        self.assertEqual(info["sheets_count"], 3)
        self.assertIn("财务统计表", info["sheets"])
        self.assertIn("Project Milestones", info["sheets"])
        self.assertIn("Empty Notes", info["sheets"])

    def test_sheet_to_html_table_merged_cells(self):
        wb = openpyxl.load_workbook(self.excel_path, data_only=True)
        ws1 = wb["财务统计表"]
        html_out, text_out = sheet_to_html_table(ws1)
        wb.close()

        # Check table tags
        self.assertIn("<table", html_out)
        self.assertIn("</table>", html_out)

        # Check colspan
        self.assertIn('colspan="2"', html_out)
        self.assertIn("部门汇总", html_out)

        # Check rowspan
        self.assertIn('rowspan="2"', html_out)
        self.assertIn("技术部", html_out)

        # Check formatted date
        self.assertIn("2026-10-21", html_out)

    def test_extract_excel_content(self):
        doc_id = "test_doc_excel"
        doc_dir = os.path.join(self.test_dir, doc_id)
        os.makedirs(doc_dir, exist_ok=True)
        shutil.copy(self.excel_path, os.path.join(doc_dir, "original.xlsx"))

        lang, chapters = extract_excel_content(doc_id, doc_dir, "test_sample.xlsx")
        self.assertEqual(len(chapters), 3)

        # Check chapter titles
        self.assertEqual(chapters[0]["title"], "财务统计表")
        self.assertEqual(chapters[1]["title"], "Project Milestones")
        self.assertEqual(chapters[2]["title"], "Empty Notes")

        # Check paragraph structure
        ch1_p = chapters[0]["paragraphs"][0]
        self.assertEqual(ch1_p["type"], "table")
        self.assertIn("<table", ch1_p["source_text"])

    async def test_service_initialization(self):
        with open(self.excel_path, "rb") as f:
            content = f.read()

        meta = await study_service.initialize_document(content, "test_sample.xlsx")
        try:
            self.assertEqual(meta["file_type"], "excel")
            self.assertEqual(meta["total_pages"], 3)
            self.assertFalse(meta["is_scanned"])
        finally:
            # Clean up created doc dir
            doc_id = meta["doc_id"]
            d_dir = study_service.get_doc_dir(doc_id)
            if os.path.exists(d_dir):
                shutil.rmtree(d_dir, ignore_errors=True)

if __name__ == "__main__":
    unittest.main()
